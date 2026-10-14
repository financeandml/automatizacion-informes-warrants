"""Libro del modelo de valoración: hojas, celdas, estilos, comentarios, nombres y tablas de datos.

El mismo objeto se calcula en Python (`excel.Evaluador`) y se guarda como .xlsx con las fórmulas vivas (openpyxl),
para que el analista lo abra en Excel y compruebe cada cifra. Convenciones de las plantillas: entradas en azul,
fórmulas en negro; importes en millones de USD y acciones en millones.
"""

import io
import re
import zipfile
from dataclasses import dataclass, field
from xml.sax.saxutils import escape

from openpyxl.utils import get_column_letter

from .excel import Error, Evaluador, Libro, TablaDeDatos

# Cálculo iterativo del libro (circularidad de intereses del LBO): más preciso que el de las plantillas (100 iteraciones,
# cambio < 0,001), cuyo residuo deja celdas del bloque circular incoherentes entre sí (p. ej. deuda inicial − amortización
# ≠ deuda final por céntimos) y hace que Python y LibreOffice paren en iteraciones distintas. Excel admite estos valores.
ITERACIONES, CAMBIO_MAXIMO = 1000, 1e-7
SIN_DATOS = "Sin datos de la SEC para este ejercicio: «--» (la fórmula de la plantilla, dentro de IFERROR)"
BURDEOS, NEGRO, BLANCO, AZUL = "2B0B02", "000000", "FFFFFF", "0000FF"
FORMATOS = {"mn": '#,##0;(#,##0);"–"', "mn1": '#,##0.0;(#,##0.0);"–"', "pct": '0.0%;(0.0%);"–"', "x": '0.0"x";(0.0"x");"–"',
            "usd": '#,##0.00;(#,##0.00)', "anio": "0", "fecha": "dd/mm/yyyy", "num": "0.00", "acc": '#,##0.0;(#,##0.0)'}


def col(n: int) -> str:
    return get_column_letter(n)


def ref(fila: int, columna: int, hoja: str | None = None, absoluta: bool = False) -> str:
    texto = f"${col(columna)}${fila}" if absoluta else f"{col(columna)}{fila}"
    return f"'{hoja}'!{texto}" if hoja else texto


@dataclass
class Celda:
    valor: object
    estilo: str = ""  # "entrada", "titulo", "seccion", "cabecera", "etiqueta", "total", ""
    formato: str | None = None
    comentario: str | None = None


@dataclass
class Hoja:
    nombre: str
    celdas: dict[tuple[int, int], Celda] = field(default_factory=dict)
    anchos: dict[int, float] = field(default_factory=dict)

    def poner(self, fila: int, columna: int, valor, estilo: str = "", formato: str | None = None, comentario: str | None = None):
        self.celdas[(fila, columna)] = Celda(valor, estilo, formato, comentario)
        return ref(fila, columna)

    def fila(self, fila: int, desde: int, valores, estilo: str = "", formato: str | None = None):
        for i, v in enumerate(valores):
            if v is not None:
                self.poner(fila, desde + i, v, estilo, formato)


@dataclass
class LibroModelo:
    hojas: list[Hoja] = field(default_factory=list)
    nombres: dict[str, str] = field(default_factory=dict)
    tablas: list[TablaDeDatos] = field(default_factory=list)
    validaciones: list[tuple[str, str, str]] = field(default_factory=list)  # (hoja, rango, lista)
    valores: dict = field(default_factory=dict)  # los del último `calcular` (se guardan junto a cada fórmula del .xlsx)

    def hoja(self, nombre: str) -> Hoja:
        h = Hoja(nombre)
        self.hojas.append(h)
        return h

    def calcular(self, max_rondas: int = 15) -> tuple[dict, list[tuple], list[tuple]]:
        """Valores del libro con los huecos de datos marcados, como hace la plantilla con «--» (p. ej. los impuestos
        negativos): una fórmula que da error SOLO porque lee una celda vacía (un ejercicio sin datos en la SEC, p. ej. una
        empresa con menos de 5 años de historia) o un «--» pasa a IFERROR(fórmula,"--"), con un comentario que lo dice.
        Devuelve (valores, celdas marcadas, celdas que siguen en error); los errores restantes no se tocan."""
        marcadas = []
        for _ in range(max_rondas):
            ev = Evaluador(self.a_libro(), iteraciones=ITERACIONES, cambio_maximo=CAMBIO_MAXIMO)
            valores = ev.calcular()
            errores = {k for k, v in valores.items() if isinstance(v, Error)}
            nuevas = []
            for k in errores:
                if k not in ev.arboles:
                    continue
                previas = ev.precedentes(k)
                if not any(p in errores for p in previas) and any(valores.get(p) in (None, "--") for p in previas):
                    nuevas.append(k)
            if not nuevas:
                self.valores = valores
                return valores, marcadas, sorted(errores)
            hojas = {h.nombre: h for h in self.hojas}
            for hoja, f, c in nuevas:
                celda = hojas[hoja].celdas[(f, c)]
                celda.valor = f'=IFERROR({celda.valor[1:]},"--")'
                celda.comentario = f"{celda.comentario} · {SIN_DATOS}" if celda.comentario else SIN_DATOS
            marcadas += nuevas
        self.valores = valores
        return valores, marcadas, sorted(errores)

    def a_libro(self) -> Libro:
        libro = Libro(hojas=[h.nombre for h in self.hojas], nombres={k.lower(): v for k, v in self.nombres.items()},
                      tablas=list(self.tablas))
        for h in self.hojas:
            for (f, c), celda in h.celdas.items():
                libro.celdas[(h.nombre, f, c)] = celda.valor
        return libro

    def guardar(self, ruta) -> None:
        """.xlsx con fórmulas vivas, estilo de Warrants & Co., cálculo iterativo (LBO) y recálculo al abrir."""
        import openpyxl
        from openpyxl.comments import Comment
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.workbook.defined_name import DefinedName
        from openpyxl.worksheet.datavalidation import DataValidation
        from openpyxl.worksheet.formula import DataTableFormula

        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        fuentes = {"titulo": Font(name="Montserrat", bold=True, size=14, color=BURDEOS),
                   "seccion": Font(name="Montserrat", bold=True, size=10, color=BLANCO),
                   "cabecera": Font(name="Inter", bold=True, size=9, color=NEGRO),
                   "etiqueta": Font(name="Inter", size=9, color=NEGRO), "total": Font(name="Inter", bold=True, size=9, color=NEGRO),
                   "entrada": Font(name="Inter", size=9, color=AZUL), "": Font(name="Inter", size=9, color=NEGRO)}
        relleno = PatternFill("solid", fgColor=BURDEOS)
        borde = Border(top=Side(style="thin", color=NEGRO))
        for h in self.hojas:
            ws = wb.create_sheet(h.nombre)
            ws.sheet_view.showGridLines = False
            for (f, c), celda in sorted(h.celdas.items()):
                x = ws.cell(row=f, column=c)
                x.value = celda.valor
                x.font = fuentes.get(celda.estilo, fuentes[""])
                if celda.estilo == "seccion":
                    x.fill = relleno
                if celda.estilo == "total":
                    x.border = borde
                if celda.formato:
                    x.number_format = FORMATOS.get(celda.formato, celda.formato)
                if celda.estilo in ("cabecera",) or (isinstance(celda.valor, (int, float)) and celda.formato):
                    x.alignment = Alignment(horizontal="right")
                if celda.comentario:
                    x.comment = Comment(celda.comentario, "Warrants & Co.")
            ws.column_dimensions["A"].width = 3
            ws.column_dimensions["B"].width = 44
            for c, ancho in h.anchos.items():
                ws.column_dimensions[col(c)].width = ancho
        for t in self.tablas:
            ws = wb[t.hoja]
            rango = f"{col(t.c1)}{t.f1}:{col(t.c2)}{t.f2}"
            r1 = ref(t.entrada_fila[1], t.entrada_fila[2]) if t.entrada_fila else None
            r2 = ref(t.entrada_columna[1], t.entrada_columna[2]) if t.entrada_columna else None
            ws.cell(row=t.f1, column=t.c1).value = DataTableFormula(rango, dt2D=bool(r1 and r2), dtr=bool(r1 and not r2),
                                                                    r1=r1 or r2, r2=r2 if r1 else None)
        for nombre, destino in self.nombres.items():
            wb.defined_names[nombre] = DefinedName(nombre, attr_text=destino)
        for hoja, rango, lista in self.validaciones:
            dv = DataValidation(type="list", formula1=f'"{lista}"', allow_blank=False)
            wb[hoja].add_data_validation(dv)
            dv.add(rango)
        wb.calculation.iterate = True  # circularidad de intereses del LBO (como las plantillas)
        wb.calculation.iterateCount = ITERACIONES
        wb.calculation.iterateDelta = CAMBIO_MAXIMO
        wb.calculation.fullCalcOnLoad = True
        buffer = io.BytesIO()
        wb.save(buffer)
        contenido = self._con_valores(buffer.getvalue()) if self.valores else buffer.getvalue()
        if hasattr(ruta, "write"):
            ruta.write(contenido)
        else:
            with open(ruta, "wb") as f:
                f.write(contenido)

    def _con_valores(self, xlsx: bytes) -> bytes:
        """Escribe junto a cada fórmula el valor calculado en Python, como hace Excel al guardar: el libro muestra las
        cifras aunque el programa que lo abra no recalcule, y LibreOffice itera desde ellas al contrastarlo (en el
        bloque circular confirma que cada fórmula se cumple, en vez de partir de cero). Excel recalcula al abrir."""
        origen, salida = zipfile.ZipFile(io.BytesIO(xlsx)), io.BytesIO()
        indices = {f"xl/worksheets/sheet{i + 1}.xml": h.nombre for i, h in enumerate(self.hojas)}
        with zipfile.ZipFile(salida, "w", zipfile.ZIP_DEFLATED) as destino:
            for item in origen.infolist():
                datos = origen.read(item.filename)
                if item.filename in indices:
                    datos = _valores_en_hoja(datos.decode("utf-8"), indices[item.filename], self.valores).encode("utf-8")
                destino.writestr(item, datos)
        return salida.getvalue()


CELDA_CON_FORMULA = re.compile(r'<c r="([A-Z]+)(\d+)"([^>]*)>(<f[^>]*/>|<f[^>]*>.*?</f>)(?:<v></v>|<v\s*/>)?</c>', re.S)


def _valores_en_hoja(xml: str, hoja: str, valores: dict) -> str:
    from openpyxl.utils import column_index_from_string

    def poner(m):
        v = valores.get((hoja, int(m.group(2)), column_index_from_string(m.group(1))))
        atributos = re.sub(r'\st="[^"]*"', "", m.group(3))
        if v is None:
            return m.group(0)
        if isinstance(v, bool):
            tipo, texto = "b", "1" if v else "0"
        elif isinstance(v, (int, float)):
            tipo, texto = None, repr(float(v))
        elif isinstance(v, Error):
            tipo, texto = "e", v.codigo
        else:
            tipo, texto = "str", escape(str(v))
        tipo = f' t="{tipo}"' if tipo else ""
        return f'<c r="{m.group(1)}{m.group(2)}"{atributos}{tipo}>{m.group(4)}<v>{texto}</v></c>'

    return CELDA_CON_FORMULA.sub(poner, xml)
